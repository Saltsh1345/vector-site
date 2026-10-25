# Список фраз страницы /v3/a/, которых нет в файле переводчика: CSV и Markdown для отправки переводчику.
import csv, re, collections, os
from openpyxl import load_workbook
D = os.path.dirname(os.path.abspath(__file__))
def norm(s):
    s = s.lower().replace('ё', 'е'); s = re.sub(r'[\[\]*_«»"“”„.,:;!?—–\-()·‖]', ' ', s); return ' '.join(s.split())
have = set()
for r in csv.DictReader(open(D + '/perevod-pervaya-versiya-155.csv', encoding='utf-8-sig')):
    have.add(norm(r['Русский']))
    for p in r['Русский'].split('‖'): have.add(norm(p))
KIND = {'Подпись для экранного диктора (не видна)': 'подпись для незрячих, не видна', 'Описание фото для экранного диктора (не видно)': 'описание фото для незрячих, не видно',
        'Описание для поиска (не видно на сайте)': 'описание для Google, не видно'}
for ru in re.findall(r'^\d+\. \*\*RU:\*\* (.*?) \*\(', open(D + '/perevod-final-94.md', encoding='utf-8').read(), re.M):
    have.add(norm(ru))
ws = load_workbook(D + '/Vector_perevody_UZ_TJ.xlsx')['Все тексты']
out, seen = [], set()
for r in range(3, ws.max_row + 1):
    ru = ws.cell(r, 4).value
    if norm(ru) in have or ru in seen: continue
    seen.add(ru); k = ws.cell(r, 3).value
    out.append([len(out) + 1, ws.cell(r, 2).value, KIND.get(k, k.lower()), ru, ws.cell(r, 5).value, ws.cell(r, 7).value])
if not out:
    for f in ('perevod-nehvataet-v3a.csv', 'perevod-nuzhno.md'):
        if os.path.exists(D + '/' + f): os.remove(D + '/' + f)
    print('не хватает: 0 — все фразы страницы переведены переводчиком'); raise SystemExit
with open(D + '/perevod-nehvataet-v3a.csv', 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f); w.writerow(['№', 'Где на сайте', 'Что это', 'Русский', 'Узбекский (латиница) — черновик студии, проверить', 'Таджикский — черновик студии, проверить']); w.writerows(out)
groups = collections.OrderedDict()
for n, sec, kind, ru, u, t in out: groups.setdefault(sec, []).append((kind, ru, u, t))
L = ['# Переводы для сайта Vector: чего не хватает', '', 'Страница-письмо: https://saltsh1345.github.io/vector-site/v3/a/', '',
     f'Ниже {len(out)} фразы и слова, которых нет в файле перевода первой версии сайта (155 строк). Нужен перевод на **узбекский (латиница)** и **таджикский**.', '',
     '- **RU** — оригинал, его переводить.',
     '- **UZ** и **TJ** — черновик студии, он сейчас стоит на сайте. Если верно — оставить, если нет — написать свой вариант.',
     '- В скобках — где стоит фраза: заголовок, кнопка, подпись на фото. «Не видно» — текст для экранного диктора (незрячих) и для поиска Google.',
     '- В черновик уже взяты слова из вашего файла: санузел — *yuvinish xonasi* / *ташноб*, консультация (UZ) — *konsultatsiya*, анонимно (TJ) — *махфӣ*, основатель (TJ) — *муассис*, офис (TJ) — *дафтар*, дежурная линия (TJ) — *телефони навбатдор*, консультант (TJ) — *мушовир*, *кумак*, *суҳбат*, *пас аз курс*.', '']
i = 0
for sec, items in groups.items():
    L += [f'## {sec}', '']
    for kind, ru, u, t in items:
        i += 1; L += [f'{i}. **RU:** {ru} *({kind})*', f'   - UZ: {u}', f'   - TJ: {t}', '']
open(D + '/perevod-nuzhno.md', 'w', encoding='utf-8', newline='\n').write('\n'.join(L).rstrip() + '\n')
print('не хватает:', len(out))
