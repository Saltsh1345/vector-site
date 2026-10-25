# Таблица для проверки переводов сайта Vector: все фразы UZ и TJ с русским оригиналом, термины, все слова.
import re, html, collections, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import os; SP = os.path.dirname(os.path.abspath(__file__))  # таблица переводов лежит рядом
OUT = sys.argv[1]
src = open(SP + '/i18n_table.py', encoding='utf-8').read()
exec(src)  # T

# раздел сайта для каждой строки таблицы — по комментариям «# ---» в файле переводов
SECT = {'шапка': 'Шапка и служебные надписи', 'первый экран': 'Первый экран (начало письма)',
        'I': 'Глава I · Вас выслушают', 'II': 'Глава II · Бесплатная консультация', 'III': 'Глава III · Если близкий против',
        'IV': 'Глава IV · Дом, а не больница', 'V': 'Глава V · Программа', 'VI–VIII': 'Главы VI–VIII · Спортзал, баня, молельная',
        'IX': 'Глава IX · Вам тоже нужна поддержка', 'X': 'Глава X · После курса', 'XI': 'Глава XI · Кто мы',
        'XII': 'Глава XII · Лицензия и документы', 'подпись': 'Подпись письма, контакты, подвал', 'кнопки': 'Кнопки (по всей странице)'}
sections, cur = [], ''
for line in src.splitlines():
    m = re.match(r'# --- (\S+)', line)
    if m:
        key = m.group(1).rstrip(',')
        key = {'первый': 'первый экран'}.get(key, key)
        cur = SECT[key]   # точное совпадение: «I» не должен ловить «II», «X» — «XII»
    if line.startswith("('"):
        sections.append(cur)
assert len(sections) == len(T), (len(sections), len(T))

BUTTONS = {'Записаться на консультацию', 'Записаться на встречу', 'Написать в Telegram', 'Позвонить</a>', 'Написать</a>', '<span>Написать</span>'}
NOTES = {
    'Комнаты на 3–4 человека': 'UZ: «hammom» здесь — санузел, а в главе VII «Hammom va hovuzcha» — баня. Различимо ли? Для бани, может, лучше «bugʻ hammomi» или «sauna». TJ: то же со словом «ҳаммом».',
    'Баня <em>и купель</em>': 'Проверить слово для бани и купели: UZ «hammom», «hovuzcha»; TJ «ҳаммом», «ҳавзча». Понятно ли, что это парная и бассейн с холодной водой?',
    'Свой санузел': 'UZ «alohida hammom», TJ «ҳаммоми алоҳида» — понятно ли, что это санузел в комнате?',
    'Мы — центр восстановления Vector, первый': 'UZ: проверить согласование «alkogol, giyohvandlik va qimorga qaram boʻlgan». TJ: «тавонбахшӣ» (реабилитация) — понятно ли, или лучше «реабилитатсия»?',
    'Основатель центра и координатор программы': 'Как основатель пишет своё имя латиницей? Сейчас: Temirxan Nurjan Kadirxanovich.',
    'alt="Дверь кабинета': 'На настоящей табличке: «Консультантская / Maslahatchi». В TJ переведено как «Машваратчӣ».',
    'content="Центр восстановления': 'Это описание страницы в поиске Google, на сайте не видно.',
}

def plain(s):
    s = re.sub(r'<b>[IVX]+</b>', '', s)                      # номер главы — не текст
    s = re.sub(r'</b><small>', ' — ', s)                     # карточка: факт — пояснение
    s = re.sub(r'<br>', ' ', s)
    m = re.match(r'^(?:<title>|aria-label=|alt=|data-cap=|content=)"?(.*?)"?(?:</title>)?$', s, re.S)
    if m and not s.startswith('<title>'):
        s = m.group(1)
    s = re.sub(r'<[^>]+>', '', s)
    s = s.replace('</a>', '')
    return html.unescape(s).strip()

def kind(s, p):
    if s.startswith('<title>'): return 'Заголовок вкладки браузера'
    if s.startswith('content='): return 'Описание для поиска (не видно на сайте)'
    if s.startswith('aria-label='): return 'Подпись для экранного диктора (не видна)'
    if s.startswith('alt='): return 'Описание фото для экранного диктора (не видно)'
    if s.startswith('data-cap='): return 'Подпись под документом в просмотре'
    if 'scene-tag' in s: return 'Подпись сцены в ролике'
    if s in BUTTONS: return 'Кнопка'
    if s.startswith('<span>Лицензия') or s.startswith('<span>Стр.') or s.startswith('<span>Сертификат'): return 'Подпись под документом'
    if '</b><small>' in s: return 'Карточка на фото: факт — пояснение'
    if s.startswith('<b>') or 'class="label"' in s: return 'Метка над заголовком'
    if s.endswith('</li>') or s.startswith('<a href="#docs">'): return 'Пометка под кнопками'
    if s.startswith('<small>'): return 'Подпись в контактах'
    if s.startswith('ООО'): return 'Подвал сайта'
    if '«Я не могу»' in s: return 'Цитата крупно'
    if len(p) < 60 and not p.endswith('.') or s.startswith('Здравствуйте'): return 'Заголовок'
    return 'Текст письма'

rows = []
for i, (ru, uz, tg) in enumerate(T):
    if ru.startswith('[[0,'):   # подписи сцен ролика — каждая отдельной строкой
        R = re.findall(r'"([^"]+)"', ru); U = re.findall(r'"([^"]+)"', uz); G = re.findall(r'"([^"]+)"', tg)
        seen = set()
        for k, (a, b, c) in enumerate(zip(R, U, G)):
            if a in seen: continue
            seen.add(a); rows.append((sections[i], 'Подпись сцены в ролике', a, b, c, '', f'T{i}.{k}'))
        continue
    p = plain(ru)
    note = next((v for k, v in NOTES.items() if ru.startswith(k)), '')
    rows.append((sections[i], kind(ru, p), p, plain(uz), plain(tg), note, f'T{i}'))
rows.insert(0, ('Шапка и служебные надписи', 'Подпись для экранного диктора (не видна)', 'Язык', 'Til', 'Забон', 'Подпись переключателя RU · UZ · TJ', 'switch'))

# ---------- книга ----------
F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF', size=10)
HFILL = PatternFill('solid', fgColor='153D36')
EDIT = PatternFill('solid', fgColor='FFF2B3')
NOTE = PatternFill('solid', fgColor='DCE8DE')
thin = Side(style='thin', color='D9D4C7')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical='top')
base = Font(name=F, size=10)

def header(ws, cols, widths, row=1):
    for j, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=j, value=c); cell.font = H; cell.fill = HFILL; cell.alignment = Alignment(wrap_text=True, vertical='center'); cell.border = B
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

wb = Workbook()
# 1. Как проверять
ws = wb.active; ws.title = 'Как проверять'
ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 92
lines = [
    ('Что это', 'Все тексты сайта центра Vector на узбекском (латиница) и таджикском, с русским оригиналом. Переводы сделаны студией и ещё не проверены носителем языка.'),
    ('Сайт', 'Русский: saltsh1345.github.io/vector-site/v3/a/ · Узбекский: …/v3/a/uz/ · Таджикский: …/v3/a/tg/'),
    ('Что заполнять', 'Только жёлтые колонки: «Правка UZ», «Правка TJ», «Комментарий». Остальное не трогать.'),
    ('Как писать правку', 'Если фраза верна — оставить пусто. Если нет — написать в жёлтой колонке исправленную фразу целиком.'),
    ('Пример правки', 'Узбекский: «Kecha-kunduz» → в «Правка UZ» написать «24 soat davomida» и в «Комментарий»: «так чаще пишут в объявлениях».'),
    ('Буквы oʻ и gʻ', 'Можно печатать как удобно: o\', g\', oʻ, o‘. При переносе на сайт приведём к правильному знаку.'),
    ('Зелёная колонка', '«Вопрос студии» — места, в которых мы сами не уверены. На них стоит посмотреть особенно внимательно.'),
    ('Листы', '«Все тексты» — каждая фраза сайта. «Термины» — главные слова, которые должны быть одинаковыми везде. «Слова UZ» и «Слова TJ» — каждое отдельное слово, которое есть на сайте.'),
    ('Где на сайте', 'Колонка «Что это»: заголовок, текст письма, кнопка, подпись на фото и т. д. Строки «не видно на сайте» читает только экранный диктор для незрячих и поиск Google.'),
    ('Сколько поправлено, UZ', "=COUNTA('Все тексты'!F3:F400)"),
    ('Сколько поправлено, TJ', "=COUNTA('Все тексты'!H3:H400)"),
    ('Сколько фраз всего', "=COUNTA('Все тексты'!D3:D400)"),
]
ws['A1'] = 'Проверка переводов · сайт Vector · 25.10.2026'; ws['A1'].font = Font(name=F, bold=True, size=14, color='153D36')
for k, (a, b) in enumerate(lines, 3):
    ws.cell(row=k, column=1, value=a).font = Font(name=F, bold=True, size=10)
    c = ws.cell(row=k, column=2, value=b); c.font = base; c.alignment = wrap
    ws.row_dimensions[k].height = 30 if len(str(b)) > 90 else 18
ws['B7'].fill = EDIT
for k in (12, 13, 14):
    ws.cell(row=k, column=2).alignment = Alignment(horizontal='left', vertical='top')
    ws.cell(row=k, column=2).font = Font(name=F, bold=True, size=11, color='153D36')

# 2. Все тексты
ws = wb.create_sheet('Все тексты')
ws['A1'] = 'Все фразы сайта. Жёлтые колонки — для правок.'; ws['A1'].font = Font(name=F, bold=True, size=11, color='153D36')
cols = ['№', 'Раздел сайта', 'Что это', 'Русский (оригинал)', 'Узбекский (латиница)', 'Правка UZ', 'Таджикский', 'Правка TJ', 'Комментарий', 'Вопрос студии', 'Код (служебно)']
header(ws, cols, [5, 24, 22, 46, 46, 30, 46, 30, 26, 34, 9], row=2)
for n, (sec, kd, ru, uz, tg, note, code) in enumerate(rows, 1):
    vals = [n, sec, kd, ru, uz, None, tg, None, None, note or None, code]
    r = n + 2
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v); c.font = base; c.alignment = wrap; c.border = B
        if j in (6, 8, 9): c.fill = EDIT
        if j == 10 and note: c.fill = NOTE
        if j == 11: c.font = Font(name=F, size=8, color='8A8A8A')
ws.auto_filter.ref = f'A2:K{len(rows) + 2}'

# 3. Термины
TERMS = [
    ('Центр восстановления', 'tiklanish markazi', 'маркази барқарорсозӣ'),
    ('Реабилитационный центр', 'reabilitatsiya markazi', 'маркази тавонбахшӣ'),
    ('Лицензия', 'litsenziya', 'иҷозатнома'),
    ('Министерство здравоохранения (Минздрав)', 'Sogʻliqni saqlash vazirligi', 'Вазорати тандурустӣ'),
    ('Стационарное лечение', 'statsionar davolash', 'табобати статсионарӣ'),
    ('Консультация', 'maslahat', 'машварат'),
    ('Бесплатно / бесплатная', 'bepul', 'ройгон'),
    ('Анонимно', 'anonim', 'беном'),
    ('Поставить на учёт', 'hisobga qoʻyish', 'ба қайд гирифтан'),
    ('Наркологический учёт', 'narkologik hisob', 'қайди наркологӣ'),
    ('Зависимость', 'qaramlik', 'вобастагӣ'),
    ('Алкоголь', 'alkogol', 'машрубот'),
    ('Наркотики', 'giyohvandlik', 'маводи мухаддир'),
    ('Азартные игры', 'qimor', 'қимор'),
    ('Специалист', 'mutaxassis', 'мутахассис'),
    ('Психолог', 'psixolog', 'равоншинос'),
    ('Клинический психолог', 'klinik psixolog', 'равоншиноси клиникӣ'),
    ('Нарколог', 'narkolog', 'нарколог'),
    ('Аддиктолог', 'addiktolog', 'аддиктолог'),
    ('Программа «12 шагов»', '«12 qadam» dasturi', 'барномаи «12 қадам»'),
    ('Близкий (человек)', 'yaqin inson, yaqiningiz', 'наздик, наздикатон'),
    ('Семья', 'oila', 'оила'),
    ('Родственники', 'qarindoshlar', 'хешовандон'),
    ('Письменное согласие', 'yozma rozilik', 'розигии хаттӣ'),
    ('Дежурная линия', 'navbatchi liniya', 'хатти навбатдор'),
    ('Круглосуточно', 'kecha-kunduz', 'шабонарӯзӣ'),
    ('Позвонить', 'qoʻngʻiroq qilish', 'занг задан'),
    ('Написать', 'yozish', 'навиштан'),
    ('Записаться', 'yozilish', 'сабти ном'),
    ('Встреча', 'uchrashuv', 'вохӯрӣ'),
    ('Суббота', 'shanba', 'шанбе'),
    ('После курса', 'kursdan keyin', 'баъди курс'),
    ('Трезвость', 'hushyorlik', 'ҳушёрӣ'),
    ('Выход (из ситуации)', 'chiqish yoʻli', 'роҳи наҷот'),
    ('Первый шаг', 'birinchi qadam', 'қадами аввал'),
    ('Комната', 'xona', 'утоқ'),
    ('Санузел', 'hammom (alohida hammom)', 'ҳаммом ва ҳоҷатхона'),
    ('Баня', 'hammom', 'ҳаммом'),
    ('Купель', 'hovuzcha', 'ҳавзча'),
    ('Спортзал', 'sport zali', 'толори варзиш'),
    ('Молельная комната', 'ibodat xonasi', 'утоқи намоз'),
    ('Двор, сад, фонтан', 'hovli, bogʻ, favvora', 'ҳавлӣ, боғ, фаввора'),
    ('Трёхразовое питание', 'kuniga uch mahal ovqat', 'се вақт хӯрок дар як рӯз'),
    ('Документы', 'hujjatlar', 'ҳуҷҷатҳо'),
    ('Основатель', 'asoschi', 'асосгузор'),
    ('Команда', 'jamoa', 'даста'),
    ('ООО', 'MChJ', 'ҶДММ'),
    ('Самарканд', 'Samarqand', 'Самарқанд'),
    ('Узбекистан', 'Oʻzbekiston', 'Ӯзбекистон'),
    ('Казахстан', 'Qozogʻiston', 'Қазоқистон'),
    ('Алматы', 'Olmaota', 'Алмаато'),
]
ws = wb.create_sheet('Термины')
ws['A1'] = 'Главные слова. Если какое-то слово неверно — оно поменяется во всех фразах, где встречается.'; ws['A1'].font = Font(name=F, bold=True, size=11, color='153D36')
header(ws, ['№', 'Русский', 'Узбекский', 'Правка UZ', 'Таджикский', 'Правка TJ', 'Комментарий'], [5, 34, 30, 26, 30, 26, 34], row=2)
for n, (a, b, c) in enumerate(TERMS, 1):
    for j, v in enumerate([n, a, b, None, c, None, None], 1):
        cell = ws.cell(row=n + 2, column=j, value=v); cell.font = base; cell.alignment = wrap; cell.border = B
        if j in (4, 6, 7): cell.fill = EDIT

# 4–5. Все слова
TJ_ABC = 'аәбвгғдеёжзиӣйкқлмнопрстуӯфхҳцчҷшщъыьэюя'
def tj_key(w):
    return [TJ_ABC.index(ch) if ch in TJ_ABC else 100 + ord(ch) for ch in w.lower()]
SKIP = {'vector', 'clinic', 'telegram', 'instagram', 'qr', 'amanat', 'rehab', 'recovery', 'center'}
def words(col, rx, keyf):
    cnt = collections.Counter(); first = {}; form = {}
    for n, r in enumerate(rows, 1):
        for w in re.findall(rx, r[col]):
            k = w.lower()
            if k in SKIP: continue
            cnt[k] += 1; first.setdefault(k, n); form.setdefault(k, k)
    return sorted(cnt.items(), key=lambda kv: keyf(kv[0])), first
UZ_RX = r"[A-Za-zʻʼ]+(?:-[A-Za-zʻʼ]+)*"
TJ_RX = r"[А-Яа-яЁёҒғӢӣҚқӮӯҲҳҶҷ]+(?:-[А-Яа-яЁёҒғӢӣҚқӮӯҲҳҶҷ]+)*"
for title, col, rx, keyf, lang in (('Слова UZ', 3, UZ_RX, lambda w: w.replace('ʻ', '').replace('ʼ', ''), 'узбекском'),
                                   ('Слова TJ', 4, TJ_RX, tj_key, 'таджикском')):
    items, first = words(col, rx, keyf)
    ws = wb.create_sheet(title)
    ws['A1'] = f'Каждое отдельное слово на {lang} с сайта ({len(items)} слов). «Сколько раз» и «Где впервые» посчитаны по текстам сайта на 25.10.2026.'
    ws['A1'].font = Font(name=F, bold=True, size=11, color='153D36')
    header(ws, ['№', 'Слово', 'Сколько раз', 'Где впервые (№ на листе «Все тексты»)', 'Правка', 'Комментарий'], [6, 30, 12, 22, 28, 36], row=2)
    for n, (w, c) in enumerate(items, 1):
        for j, v in enumerate([n, w, c, first[w], None, None], 1):
            cell = ws.cell(row=n + 2, column=j, value=v); cell.font = base; cell.border = B
            if j in (5, 6): cell.fill = EDIT
    ws.auto_filter.ref = f'A2:F{len(items) + 2}'
    print(title, len(items))

for s in wb.worksheets:
    s.sheet_view.zoomScale = 110
wb.save(OUT)
print('фраз', len(rows), 'терминов', len(TERMS))
